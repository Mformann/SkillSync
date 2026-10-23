"""Phase 4: factual views, shared validation, layout-aware bounded extraction."""
import io
import json

from docx import Document
from docx.oxml import OxmlElement
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session
import pytest

from app import models
from app.services.ai_analysis_service import AIAnalysis, Requirement, RequirementMatch, Strength
from app.services.document_parser import parse_resume, ResumeParseError, ParserLimits
from test_analysis_runs import environment, provider, duplicate_analysis, endpoint, JOB, RESUME

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def stats(client):
    response = client.get("/dashboard/summary")
    response.raise_for_status()
    return {item["label"]: item["value"] for item in response.json()["stats"]}


def test_zero_is_preserved_in_every_current_view(environment, provider):
    client, factory, aid, wid, *_ = environment
    client.patch(f"/workspaces/{wid}", json={"description": "We require Kubernetes expertise for reliable production application delivery."}).raise_for_status()
    assert client.post(endpoint(aid)).json()["scores"]["overall"] == 0
    with factory() as db:
        analysis = db.get(models.Analysis, aid)
        analysis.analysis_data = {"overview": {"overallScore": 78}}
        analysis.score_breakdown = {"overall": 78}
        db.commit()
    assert client.get("/resume/all").json()[0]["score"] == 0
    assert stats(client) == {"Latest Score": 0, "Total Resumes": 1, "Avg Score": 0, "Analysis Done": 1}
    assert client.get(f"/resume/analysis/{aid}").json()["overview"]["overallScore"] == 0
    assert client.get(f"/resume/analysis/{aid}/explainable").json()["scores"]["overall"] == 0
    career = client.get("/career/workspaces").json()[0]["readiness"]
    assert career["components"]["job_match"] == 0
    for path in (f"/resume/gap-report/{aid}", "/resume/gap-report/latest"):
        report = client.get(path).json()
        assert report["targetRoles"][0]["matchScore"] == 0
        assert report["projection_version"] == "evidence-views-v1"


def test_gaps_goals_and_strengths_are_job_specific_and_evidenced(environment, provider):
    client, factory, aid, *_ = environment
    client.post(endpoint(aid)).raise_for_status()
    report = client.get(f"/resume/gap-report/{aid}").json()
    assert [gap["skill"] for gap in report["gapAnalysis"]] == ["Docker"]
    assert report["strengthsVsGaps"]["strengths"][0]["skill"] == "Python"
    assert report["strengthsVsGaps"]["strengths"][0]["evidence"][0] in RESUME
    for forbidden in ("Expert", '"years"', "timeToReady", "demandTrend", "Azure", "AWS"):
        assert forbidden not in json.dumps(report)
    summary = client.get("/dashboard/summary").json()
    assert [goal["skill"] for goal in summary["weeklyGoals"]] == ["Docker"]
    assert summary["weeklyGoals"][0]["goal"].startswith("Docker:")
    assert summary["achievements"] == []
    assert len(summary["progressData"]) == 1
    assert summary["progressData"][0]["score"] == report["targetRoles"][0]["matchScore"]


def test_pending_failed_and_processing_do_not_become_zero_or_completed(environment, provider):
    client, factory, aid, *_ = environment
    client.post(endpoint(aid)).raise_for_status()
    other = duplicate_analysis(factory, aid)
    assert client.get("/resume/all").json()[0]["score"] is None
    assert stats(client) == {"Latest Score": None, "Total Resumes": 2, "Avg Score": 50, "Analysis Done": 1}
    career = client.get("/career/workspaces").json()[0]["readiness"]
    assert career["components"]["job_match"] is None and career["overall"] is None
    for status in ("processing", "failed"):
        with factory() as db:
            db.get(models.Analysis, aid).analysis_status = status
            db.commit()
        assert stats(client)["Analysis Done"] == 0
        assert stats(client)["Avg Score"] is None
        assert client.get(f"/resume/analysis/{aid}/explainable").json()["scores"] is None
    assert client.get(f"/resume/analysis/{other}").json()["status"] == "pending"


def test_stale_and_legacy_sample_data_are_hidden_without_rewriting_history(environment, provider):
    client, factory, aid, wid, *_ = environment
    client.post(endpoint(aid)).raise_for_status()
    client.patch(f"/workspaces/{wid}", json={"description": "We require Kubernetes expertise for reliable production application delivery."}).raise_for_status()
    assert client.get("/resume/all").json()[0]["score"] is None
    assert client.get(f"/resume/gap-report/{aid}").json()["status"] == "stale"
    assert client.get("/dashboard/summary").json()["weeklyGoals"] == []
    with factory() as db:
        analysis = db.get(models.Analysis, aid)
        analysis.analysis_data = {"overview": {"overallScore": 78}}
        analysis.gap_report = {"fiction": "Expert 3+ years"}
        analysis.explainable_data = None
        analysis.input_hash = None
        db.commit()
    for path in (f"/resume/analysis/{aid}", "/resume/analysis/latest", f"/resume/gap-report/{aid}",
                 "/resume/gap-report/latest", f"/resume/roadmap/{aid}", "/resume/roadmap/latest"):
        response = client.get(path)
        assert response.status_code == 200, response.text
        assert response.json()["status"] == "legacy"
        assert "Expert" not in response.text and "78" not in response.text
    with factory() as db:
        assert db.get(models.Analysis, aid).analysis_data["overview"]["overallScore"] == 78
        assert db.get(models.Analysis, aid).gap_report == {"fiction": "Expert 3+ years"}


def test_no_detected_requirements_means_unavailable_not_zero(environment, provider):
    client, _, aid, wid, *_ = environment
    client.patch(f"/workspaces/{wid}", json={"description": "Join our team to deliver excellent customer outcomes and organizational initiatives."}).raise_for_status()
    response = client.post(endpoint(aid)).json()
    assert response["status"] == "unavailable" and response["scores"] is None
    assert stats(client)["Latest Score"] is None
    assert stats(client)["Analysis Done"] == 0


def test_invalid_legacy_result_cannot_be_reused_from_cache(environment, provider):
    client, factory, aid, *_ = environment
    client.post(endpoint(aid)).raise_for_status()
    with factory() as db:
        db.get(models.Analysis, aid).explainable_data = {"unsupported_sample": 78}
        db.commit()
    assert client.get(f"/resume/analysis/{aid}/explainable").json()["status"] == "legacy"
    other = duplicate_analysis(factory, aid)
    response = client.post(endpoint(other))
    assert response.status_code == 200 and response.json()["status"] == "completed_ai"
    assert len(provider) == 2
    response = client.post(endpoint(aid))
    assert response.status_code == 200 and response.json()["status"] == "completed_cached"
    assert response.json()["scores"]["overall"] == 50


def test_priority_and_supported_strength_filtering(environment, monkeypatch):
    client, _, aid, *_ = environment
    async def fake(*_):
        return AIAnalysis(summary="Synthetic", candidate_positioning="Synthetic", requirements=[
            Requirement(name=name, category="technical_skill", priority=priority, importance=importance, rationale="In job")
            for name, priority, importance in (("React", "preferred", 5), ("Python", "required", 1), ("Docker", "required", 5))],
            matches=[RequirementMatch(requirement_name=name, status="missing_skill", evidence_quotes=[], explanation="Missing", confidence=0.8, recommended_action=f"Demonstrate {name}.") for name in ("React", "Python", "Docker")],
            strengths=[Strength(title="Invented expert", evidence="3+ years of unsupported expertise", why_relevant="Invented")], resume_risks=[])
    monkeypatch.setattr("app.services.ai_analysis_service.GroqProvider.analyze_async", fake)
    result = client.post(endpoint(aid)).json()
    assert result["result"]["strengths"] == []
    assert [gap["skill"] for gap in client.get(f"/resume/gap-report/{aid}").json()["gapAnalysis"]] == ["Docker", "Python", "React"]


@pytest.fixture
def fast_parser(monkeypatch):
    calls = []
    def parse(*_):
        calls.append(True)
        return RESUME
    monkeypatch.setattr("app.routers.resume.parse_resume", parse)
    return calls


def upload(client, fields=None, *, wid=None, content=b"PKsynthetic"):
    data = {"job_role": "Python engineer", "job_description": JOB, "company": "Example", "ai_consent": "false"}
    data.update(fields or {})
    if wid is not None:
        data["workspace_id"] = str(wid)
    return client.post("/resume/upload", data=data, files={"file": ("synthetic.docx", content, DOCX)})


@pytest.mark.parametrize("field,value", [("title", " "), ("title", "A"), ("title", "x"*161),
    ("description", " "*70), ("description", "x"*49), ("description", "x"*50001), ("company", "x"*161)],
    ids=["blank-title", "short-title", "long-title", "blank-description", "short-description", "long-description", "long-company"])
def test_json_and_upload_share_trimmed_limits(environment, fast_parser, field, value):
    client, factory, *_ = environment
    response = client.post("/workspaces", json={"title": "Python engineer", "description": JOB, **{field: value}})
    assert response.status_code == 422 and isinstance(response.json()["detail"], list)
    source = {"title": "job_role", "description": "job_description", "company": "company"}[field]
    response = upload(client, {source: value})
    assert response.status_code == 422 and response.json()["detail"][0]["loc"][-1] == source
    assert fast_parser == []
    with factory() as db:
        assert db.query(models.TargetJob).count() == db.query(models.Resume).count() == db.query(models.Analysis).count() == 1
        assert db.query(models.AIUsage).count() == 0


@pytest.mark.parametrize("field", ["title", "description", "ai_consent", "is_archived"])
def test_explicit_null_is_rejected_but_omitted_fields_are_preserved(environment, field):
    client, factory, _, wid, *_ = environment
    response = client.patch(f"/workspaces/{wid}", json={field: None})
    assert response.status_code == 422 and response.json()["detail"][0]["loc"][-1] == field
    assert client.patch(f"/workspaces/{wid}", json={}).status_code == 200
    with factory() as db:
        job = db.get(models.TargetJob, wid)
        assert job.title == "Python engineer" and job.description == JOB and job.ai_consent is True and job.is_archived is False


def test_optional_clear_and_maximum_valid_values_succeed(environment, fast_parser):
    client, factory, _, wid, *_ = environment
    values = {"title": " " + "T"*160 + " ", "company": " " + "C"*160 + " ", "description": " " + "D"*50000 + " "}
    created = client.post("/workspaces", json=values)
    assert created.status_code == 201 and len(created.json()["title"]) == 160
    uploaded = upload(client, {"job_role": values["title"], "company": values["company"], "job_description": values["description"]})
    assert uploaded.status_code == 200, uploaded.text
    response = client.patch(f"/workspaces/{wid}", json={"company": None, "location": " ", "employment_type": None})
    assert response.status_code == 200
    assert response.json()["company"] is response.json()["location"] is response.json()["employment_type"] is None


def test_saved_job_upload_is_atomic_and_does_not_duplicate_workspaces(environment, fast_parser, monkeypatch):
    client, factory, _, wid, *_ = environment
    def rejected(*_):
        raise ResumeParseError("invalid")
    monkeypatch.setattr("app.routers.resume.parse_resume", rejected)
    response = upload(client, {"job_role": "Changed role", "job_description": JOB + " Updated.", "ai_consent": "false"}, wid=wid)
    assert response.status_code == 422
    with factory() as db:
        assert db.get(models.TargetJob, wid).title == "Python engineer"
        assert db.get(models.TargetJob, wid).ai_consent is True
    monkeypatch.setattr("app.routers.resume.parse_resume", lambda *_: RESUME)
    response = upload(client, {"job_role": " Changed role ", "company": "", "ai_consent": "false"}, wid=wid)
    assert response.status_code == 200 and response.json()["workspace_id"] == wid
    with factory() as db:
        assert db.query(models.TargetJob).count() == 1
        assert db.get(models.TargetJob, wid).title == "Changed role" and db.get(models.TargetJob, wid).company is None
        assert db.get(models.TargetJob, wid).ai_consent is False
        assert db.query(models.Resume).count() == db.query(models.Analysis).count() == 2


@pytest.mark.parametrize("saved", [False, True])
def test_failed_save_rolls_back_all_upload_writes(environment, fast_parser, monkeypatch, saved):
    client, factory, _, wid, *_ = environment
    monkeypatch.setattr(Session, "commit", lambda self: (_ for _ in ()).throw(SQLAlchemyError("Synthetic save failure")))
    response = upload(client, {"job_role": "Changed role"}, wid=wid if saved else None)
    assert response.status_code == 503
    with factory() as db:
        assert db.get(models.TargetJob, wid).title == "Python engineer"
        for model in (models.TargetJob, models.Resume, models.Analysis):
            assert db.query(model).count() == 1
        assert db.query(models.AIUsageAccount).count() == db.query(models.AIUsage).count() == 0


def document_bytes(document):
    output = io.BytesIO(); document.save(output); return output.getvalue()


def test_docx_mixed_nested_and_merged_cells_are_in_order_once():
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Merged left"
    table.cell(0, 1).text = "Merged right"
    table.cell(0, 0).merge(table.cell(0, 1))
    cell = table.cell(1, 0)
    cell.text = "Before nested"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Nested evidence"
    cell.add_paragraph("After nested")
    table.cell(1, 1).text = "Final cell"
    vertical = document.add_table(rows=2, cols=1)
    vertical.cell(0, 0).text = "Vertical first"
    vertical.cell(1, 0).text = "Vertical second"
    vertical.cell(0, 0).merge(vertical.cell(1, 0))
    document.add_paragraph("After table")
    text = parse_resume(document_bytes(document), DOCX)
    assert [line for line in text.splitlines() if line] == ["Before table", "Merged left", "Merged right", "Before nested", "Nested evidence", "After nested", "Final cell", "Vertical first", "Vertical second", "After table"]


def test_unique_headers_and_footers_are_read_once_and_images_are_not_ocr():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Header contact"
    document.add_paragraph("Body evidence")
    document.sections[0].footer.paragraphs[0].text = "Footer link"
    document.add_section()
    drawing = OxmlElement("w:drawing")
    box = OxmlElement("w:txbxContent"); paragraph = OxmlElement("w:p"); run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = "Unsupported drawing text"
    run.append(text); paragraph.append(run); box.append(paragraph); drawing.append(box)
    document.add_paragraph().add_run()._r.append(drawing)
    text = parse_resume(document_bytes(document), DOCX)
    assert [line for line in text.splitlines() if line] == ["Header contact", "Body evidence", "Footer link"]


def test_table_text_is_subject_to_the_same_output_limit():
    document = Document(); document.add_table(rows=1, cols=1).cell(0, 0).text = "x"*500
    with pytest.raises(ResumeParseError) as error:
        parse_resume(document_bytes(document), DOCX, ParserLimits(max_text_chars=100))
    assert error.value.code == "limit"
