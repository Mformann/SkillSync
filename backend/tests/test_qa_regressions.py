"""QA probes: synthetic data and isolated SQLite; no external AI or account writes."""
import io
import base64
import json
from datetime import date, timedelta

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import models
from app.auth_utils import AuthenticatedUser, get_current_user
from app.database import Base, get_session
from app.main import create_app
from app.services.ai_analysis_service import ProviderResult, RuleBasedProvider

JD = "We require Python and Docker to build reliable backend APIs and automated tests."
RESUME = "QA Candidate\nBuilt production Python APIs and SQL services with thorough automated tests."


@pytest.fixture
def env(monkeypatch):
    monkeypatch.setenv("ENVIRONMENT", "development")
    monkeypatch.setenv("GROQ_API_KEY", "")
    # Lifespan must never initialize the database configured in a local .env.
    monkeypatch.setattr("app.main.init_db", lambda: None)
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    @event.listens_for(engine, "connect")
    def foreign_keys(connection, _):
        connection.execute("PRAGMA foreign_keys=ON")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    app = create_app()
    identity = {"id": "qa-user-a"}
    def database():
        with factory() as db:
            yield db
    app.dependency_overrides[get_session] = database
    app.dependency_overrides[get_current_user] = lambda: AuthenticatedUser(id=identity["id"])
    with TestClient(app, raise_server_exceptions=False) as client:
        yield client, factory, identity
    engine.dispose()


def docx_bytes(table=False):
    document = Document()
    if table:
        document.add_table(rows=1, cols=1).cell(0, 0).text = RESUME
    else:
        document.add_paragraph(RESUME)
    content = io.BytesIO()
    document.save(content)
    return content.getvalue()


def upload(client, *, consent=False, title="QA engineer", description=JD, table=False):
    return client.post("/resume/upload", data={"job_role": title, "job_description": description,
        "ai_consent": str(consent).lower()}, files={"file": ("qa-resume.docx", docx_bytes(table),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})


def test_full_candidate_flow_and_ownership(env):
    client, factory, identity = env
    response = upload(client)
    assert response.status_code == 200, response.text
    aid, wid = response.json()["analysis_id"], response.json()["workspace_id"]
    result = client.post(f"/resume/analysis/{aid}/run")
    assert result.status_code == 200, result.text
    assert result.json()["provider"] == "rule_based"
    plan = client.post("/learning/plans", json={"analysis_id": aid, "hours_per_week": 5,
        "target_date": str(date.today() + timedelta(days=60)), "experience_level": "intermediate"})
    assert plan.status_code == 201, plan.text
    version = client.post(f"/resume-studio/analyses/{aid}/versions", json={})
    assert version.status_code == 201, version.text
    vid = version.json()["id"]
    assert client.get(f"/resume-studio/versions/{vid}/docx").content.startswith(b"PK")
    exported = client.get("/product/privacy/export")
    assert exported.status_code == 200 and len(exported.json()["resumes"]) == 1
    identity["id"] = "qa-user-b"
    for path in [f"/workspaces/{wid}", f"/resume/analysis/{aid}/explainable",
                 f"/learning/plans/by-analysis/{aid}", f"/resume-studio/versions/{vid}/docx"]:
        assert client.get(path).status_code == 404, path
    assert client.get("/product/privacy/export").json()["resumes"] == []
    identity["id"] = "qa-user-a"
    deleted = client.post("/product/privacy/delete-app-data", json={"confirmation": "DELETE MY SKILLSYNC DATA"})
    assert deleted.status_code == 200, deleted.text
    assert client.get("/workspaces").json() == []
    assert client.get("/resume/all").json() == []


def test_zero_match_is_not_replaced_by_sample_score(env):
    client, _, _ = env
    response = upload(client, description="We require Kubernetes and Docker expertise to operate scalable production platforms.")
    aid = response.json()["analysis_id"]
    result = client.post(f"/resume/analysis/{aid}/run").json()
    assert result["scores"]["overall"] == 0
    assert client.get("/resume/all").json()[0]["score"] == 0


def test_daily_ai_limit_counts_forced_reruns(env, monkeypatch):
    client, _, _ = env
    calls = []
    async def fake_ai(self, job, resume):
        calls.append(True)
        return RuleBasedProvider().analyze(job, resume)
    monkeypatch.setattr("app.services.ai_analysis_service.GroqProvider.analyze_async", fake_ai)
    monkeypatch.setenv("GROQ_API_KEY", "synthetic-qa-key")
    monkeypatch.setenv("AI_DAILY_ANALYSIS_LIMIT", "1")
    aid = upload(client, consent=True).json()["analysis_id"]
    assert client.post(f"/resume/analysis/{aid}/run").status_code == 200
    assert client.post(f"/resume/analysis/{aid}/run?force=true").status_code == 200
    assert len(calls) == 1, "A forced rerun bypassed the one-call daily allowance"


def test_updated_job_invalidates_analysis(env):
    client, _, _ = env
    data = upload(client).json()
    aid, wid = data["analysis_id"], data["workspace_id"]
    first = client.post(f"/resume/analysis/{aid}/run").json()
    assert any(r["name"] == "Python" for r in first["result"]["requirements"])
    patched = client.patch(f"/workspaces/{wid}", json={"description": "We require Kubernetes expertise to run production distributed systems reliably."})
    assert patched.status_code == 200
    second = client.post(f"/resume/analysis/{aid}/run").json()
    assert [r["name"] for r in second["result"]["requirements"]] == ["Kubernetes"]


def test_table_resume_is_readable(env):
    client, _, _ = env
    response = upload(client, table=True)
    assert response.status_code == 200, response.text


def test_upload_enforces_workspace_field_limits(env):
    client, _, _ = env
    response = upload(client, title="Q" * 161)
    assert response.status_code == 422, response.text


def test_null_workspace_title_is_validation_error(env):
    client, _, _ = env
    created = client.post("/workspaces", json={"title": "QA role", "description": JD}).json()
    response = client.patch(f"/workspaces/{created['id']}", json={"title": None})
    assert response.status_code == 422, response.text


def test_deeply_nested_unsigned_token_is_unauthorized(monkeypatch):
    monkeypatch.setenv("SUPABASE_URL", "https://qa-fixture.supabase.co")
    monkeypatch.setenv("SUPABASE_PUBLISHABLE_KEY", "sb_publishable_fixture")
    def b64(value):
        return base64.urlsafe_b64encode(value).rstrip(b"=").decode()
    header = b64(json.dumps({"alg": "HS256", "typ": "JWT"}).encode())
    payload = b64(b"[" * 6000 + b"]" * 6000)
    token = header + "." + payload + "." + b64(b"synthetic-forged-signature")
    assert len(token) < 16384
    client = TestClient(create_app(), raise_server_exceptions=False)
    response = client.get("/auth/me", headers={"Authorization": "Bearer " + token})
    assert response.status_code == 401, response.text


def test_workspace_listing_does_not_query_each_row_separately(env):
    client, factory, _ = env
    with factory() as db:
        db.add_all(models.TargetJob(user_id="qa-user-a", title=f"QA role {i}", description=JD) for i in range(20))
        db.commit()
    statements = []
    engine = factory.kw["bind"]
    def record(_conn, _cursor, statement, *_):
        if statement.lstrip().upper().startswith("SELECT"):
            statements.append(statement)
    event.listen(engine, "before_cursor_execute", record)
    try:
        response = client.get("/workspaces")
        assert response.status_code == 200
    finally:
        event.remove(engine, "before_cursor_execute", record)
    assert len(statements) <= 5, f"Listing 20 workspaces executed {len(statements)} SELECT queries"


def test_gap_report_uses_target_job_requirements(env):
    client, _, _ = env
    data = upload(client).json()
    aid = data["analysis_id"]
    client.post(f"/resume/analysis/{aid}/run").raise_for_status()
    report = client.get(f"/resume/gap-report/{aid}")
    assert report.status_code == 200, report.text
    gaps = {item["skill"].casefold() for item in report.json()["gapAnalysis"]}
    assert gaps == {"docker"}, f"The target requires Python and Docker, but gap report returned {sorted(gaps)}"
