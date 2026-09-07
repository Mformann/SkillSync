from __future__ import annotations

import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

SUPPORTED_STATUSES = {"strong_match", "partial_match", "transferable", "poorly_demonstrated"}
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d[\d ().-]{7,}\d)")
URL_RE = re.compile(r"(?:https?://|www\.|linkedin\.com/|github\.com/)\S+", re.I)
NUMBER_RE = re.compile(r"(?<!\w)\d+(?:[.,]\d+)?%?(?!\w)")


def _contact(resume_text: str) -> dict:
    lines = [line.strip() for line in resume_text.splitlines() if line.strip()]
    email = EMAIL_RE.search(resume_text)
    phone = PHONE_RE.search(resume_text)
    links = list(dict.fromkeys(match.rstrip(".,)") for match in URL_RE.findall(resume_text)))
    name = ""
    for line in lines[:8]:
        if len(line) <= 80 and not EMAIL_RE.search(line) and not PHONE_RE.search(line) and not URL_RE.search(line):
            name = line
            break
    return {
        "name": name,
        "email": email.group(0) if email else "",
        "phone": phone.group(0) if phone else "",
        "location": "",
        "links": links[:4],
    }


def build_grounded_draft(analysis) -> tuple[dict, dict]:
    data = analysis.explainable_data or {}
    matches = data.get("matches", [])
    supported = [item for item in matches if item.get("status") in SUPPORTED_STATUSES]
    skills = list(dict.fromkeys(item.get("requirement_name", "").strip() for item in supported if item.get("requirement_name")))
    evidence = []
    source_map = {}
    for match in supported:
        for quote in match.get("evidence_quotes", [])[:2]:
            quote = quote.strip()
            if not quote:
                continue
            key = f"evidence-{len(evidence) + 1}"
            evidence.append({"id": key, "text": quote, "requirement": match.get("requirement_name", "")})
            source_map[key] = {"quote": quote, "requirement": match.get("requirement_name", "")}
    role = analysis.target_job.title if analysis.target_job else (analysis.job_role or "Target role")
    summary = (
        f"Candidate targeting {role} with verified resume evidence aligned to "
        f"{', '.join(skills[:5])}."
        if skills else f"Candidate targeting {role}. Add a factual professional summary before exporting."
    )
    content = {
        "basics": _contact(analysis.resume_text or (analysis.resume.extracted_text if analysis.resume else "")),
        "headline": role,
        "summary": summary,
        "skills": skills,
        "evidence": evidence,
    }
    source_map["verified_skills"] = skills
    source_map["summary_requirements"] = skills[:5]
    return content, source_map


def claim_warnings(content: dict, source_map: dict, fallback: bool = False) -> list[dict]:
    warnings = []
    basics = content.get("basics", {})
    if not basics.get("name"):
        warnings.append({"code": "missing_name", "severity": "high", "message": "Add your name before exporting."})
    if not basics.get("email") and not basics.get("phone"):
        warnings.append({"code": "missing_contact", "severity": "high", "message": "Add at least one contact method."})
    verified = {str(item).casefold() for item in source_map.get("verified_skills", [])}
    added = [item for item in content.get("skills", []) if str(item).casefold() not in verified]
    if added:
        warnings.append({
            "code": "unverified_skills", "severity": "medium",
            "message": f"Confirm manually added skills: {', '.join(added[:6])}.",
        })
    for item in content.get("evidence", []):
        source = source_map.get(item.get("id", ""), {}).get("quote", "")
        new_numbers = set(NUMBER_RE.findall(item.get("text", ""))) - set(NUMBER_RE.findall(source))
        if new_numbers:
            warnings.append({
                "code": "unsupported_metric", "severity": "high",
                "evidence_id": item.get("id"),
                "message": f"Verify newly added metric(s): {', '.join(sorted(new_numbers))}.",
            })
    if fallback:
        warnings.append({
            "code": "limited_analysis", "severity": "medium",
            "message": "The source analysis used limited keyword matching; review every tailored section.",
        })
    return warnings


def export_docx(content: dict, template: str) -> BytesIO:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.72)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    basics = content.get("basics", {})
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER if template == "classic" else WD_ALIGN_PARAGRAPH.LEFT
    run = heading.add_run(basics.get("name") or "YOUR NAME")
    run.bold = True
    run.font.size = Pt(20)
    contacts = [basics.get(key) for key in ("email", "phone", "location") if basics.get(key)]
    contacts.extend(basics.get("links", []))
    if contacts:
        p = document.add_paragraph(" | ".join(contacts))
        p.alignment = heading.alignment
    if content.get("headline"):
        p = document.add_paragraph()
        p.alignment = heading.alignment
        p.add_run(content["headline"]).bold = True
    if content.get("summary"):
        document.add_heading("Professional Summary", level=1)
        document.add_paragraph(content["summary"])
    if content.get("skills"):
        document.add_heading("Relevant Skills", level=1)
        document.add_paragraph(" • ".join(content["skills"]))
    if content.get("evidence"):
        document.add_heading("Selected Evidence", level=1)
        for item in content["evidence"]:
            if item.get("text"):
                document.add_paragraph(item["text"], style="List Bullet")
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
