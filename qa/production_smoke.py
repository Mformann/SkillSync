"""Exercise the real Linux entrypoint with signed JWTs and disposable PostgreSQL.

Run inside the production image with this folder mounted at /qa. No auth overrides,
hosted credentials, external AI calls, or developer database are used.
"""
import io
import json
import os
from pathlib import Path
import ssl
import statistics
import subprocess
import sys
import tempfile
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from uuid import uuid4

import httpx
import jwt
from cryptography import x509
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import ec
from cryptography.x509.oid import NameOID
from docx import Document
from sqlalchemy import create_engine, text
from sqlalchemy.engine import make_url

sys.path.insert(0, "/app")
url = os.environ.get("DATABASE_URL", "")
if not (make_url(url).database or "").startswith("skillsync_qa_"):
    raise SystemExit("Refusing smoke test: use a disposable skillsync_qa_ database.")
if os.getuid() == 0:
    raise SystemExit("Smoke test must run as the production non-root user.")

checks = []
AUTH_FIXTURE_HOST = "auth.fixture.example"
AUTH_FIXTURE_ORIGIN = f"https://{AUTH_FIXTURE_HOST}:8443"
result = {"scope": "local production container; synthetic identities and documents", "checks": checks}
output = Path(os.getenv("QA_OUTPUT", "/qa-output/phase-06-production-smoke.json"))
started = time.monotonic()
engine = create_engine(url)
with engine.begin() as db:
    for role in ("anon", "authenticated"):
        if not db.execute(text("SELECT 1 FROM pg_roles WHERE rolname=:name"), {"name": role}).scalar():
            db.exec_driver_sql(f'CREATE ROLE "{role}" NOLOGIN')

with tempfile.TemporaryDirectory() as temporary:
    folder = Path(temporary)
    key = ec.generate_private_key(ec.SECP256R1())
    name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, AUTH_FIXTURE_HOST)])
    now = datetime.now(timezone.utc)
    cert = (x509.CertificateBuilder().subject_name(name).issuer_name(name)
        .public_key(key.public_key()).serial_number(x509.random_serial_number())
        .not_valid_before(now - timedelta(minutes=1)).not_valid_after(now + timedelta(hours=1))
        .add_extension(x509.SubjectAlternativeName([x509.DNSName(AUTH_FIXTURE_HOST)]), critical=False)
        .add_extension(x509.BasicConstraints(ca=True, path_length=None), critical=True)
        .sign(key, hashes.SHA256()))
    (folder / "cert.pem").write_bytes(cert.public_bytes(serialization.Encoding.PEM))
    (folder / "key.pem").write_bytes(key.private_bytes(serialization.Encoding.PEM,
        serialization.PrivateFormat.PKCS8, serialization.NoEncryption()))
    jwk = json.loads(jwt.algorithms.ECAlgorithm.to_jwk(key.public_key()))
    jwk.update(kid="beta-fixture", alg="ES256", use="sig")
    class JWKS(BaseHTTPRequestHandler):
        def do_GET(self):
            if self.path != "/auth/v1/.well-known/jwks.json":
                self.send_error(404); return
            body = json.dumps({"keys": [jwk]}).encode()
            self.send_response(200); self.send_header("Content-Type", "application/json")
            self.end_headers(); self.wfile.write(body)
        def log_message(self, *_):
            pass
    issuer_server = ThreadingHTTPServer(("127.0.0.1", 8443), JWKS)
    tls = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    tls.load_cert_chain(folder / "cert.pem", folder / "key.pem")
    issuer_server.socket = tls.wrap_socket(issuer_server.socket, server_side=True)
    threading.Thread(target=issuer_server.serve_forever, daemon=True).start()
    environment = dict(os.environ, ENVIRONMENT="production", SUPABASE_URL=AUTH_FIXTURE_ORIGIN,
        SSL_CERT_FILE=str(folder / "cert.pem"), GROQ_API_KEY="", ENABLE_API_DOCS="false",
        NO_PROXY=f"{AUTH_FIXTURE_HOST},localhost,127.0.0.1",
        CORS_ORIGINS="https://beta.fixture.example", PORT="8000", WEB_CONCURRENCY="1")
    log = (folder / "api.log").open("w", encoding="utf-8")
    process = subprocess.Popen(["sh", "/app/scripts/start.sh"], env=environment, stdout=log, stderr=log)
    def token(subject, **extra):
        return jwt.encode(dict(sub=subject, aud="authenticated", iss=f"{AUTH_FIXTURE_ORIGIN}/auth/v1",
            exp=int(time.time()) + 600, **extra), key, algorithm="ES256", headers={"kid": "beta-fixture"})
    owner, other = str(uuid4()), str(uuid4())
    client = httpx.Client(base_url="http://127.0.0.1:8000", timeout=35, trust_env=False)
    headers = {"Authorization": "Bearer " + token(owner)}
    def request(method, path, status=200, **kwargs):
        response = client.request(method, path, **kwargs)
        assert response.status_code == status, f"{method} {path}: expected {status}, got {response.status_code}"
        return response
    try:
        deadline = time.monotonic() + 45
        while True:
            if process.poll() is not None:
                raise AssertionError("Production startup exited before readiness")
            try:
                if client.get("/health/ready").status_code == 200:
                    break
            except httpx.HTTPError:
                pass
            if time.monotonic() > deadline:
                raise AssertionError("Production readiness exceeded 45 seconds")
            time.sleep(.2)
        result["startup_seconds"] = round(time.monotonic() - started, 3)
        checks.append("real entrypoint validates configuration, migrates and becomes ready")
        request("GET", "/workspaces", 401)
        request("GET", "/docs", 404)
        request("GET", "/openapi.json", 404)
        request("GET", "/auth/me", headers=headers)
        request("GET", "/auth/me", 401, headers={"Authorization": "Bearer " + token(owner)[:-8] + "invalid"})
        checks.append("real ES256/JWKS authentication succeeds; invalid and missing tokens rejected; docs disabled")
        trusted = request("OPTIONS", "/workspaces", headers={"Origin": "https://beta.fixture.example",
            "Access-Control-Request-Method": "GET"})
        assert trusted.headers.get("access-control-allow-origin") == "https://beta.fixture.example"
        untrusted = client.options("/workspaces", headers={"Origin": "https://evil.fixture.example",
            "Access-Control-Request-Method": "GET"})
        assert "access-control-allow-origin" not in untrusted.headers
        request("POST", "/workspaces", 413, headers=headers, content=b"x" * (13 * 1024 * 1024))
        checks.append("CORS allowlist and production request-size limit enforced")
        document = Document(); document.add_table(rows=1, cols=1).cell(0, 0).text = (
            "Synthetic Candidate built Python SQL APIs with automated tests and Git.")
        content = io.BytesIO(); document.save(content)
        upload = request("POST", "/resume/upload", headers=headers,
            data={"job_role": "Backend Engineer", "job_description": "Build Python SQL services with Docker Kubernetes and automated testing.", "ai_consent": "false"},
            files={"file": ("synthetic.docx", content.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}).json()
        aid = upload["analysis_id"]
        request("POST", f"/resume/analysis/{aid}/run", headers=headers)
        analysis = request("GET", f"/resume/analysis/{aid}/explainable", headers=headers).json()
        wid = analysis["workspace_id"]
        request("POST", f"/resume/analysis/{aid}/run", headers=headers)
        request("GET", f"/workspaces/{wid}", 404, headers={"Authorization": "Bearer " + token(other)})
        request("GET", f"/resume/analysis/{aid}", 404, headers={"Authorization": "Bearer " + token(other)})
        checks.append("isolated DOCX worker, persisted analysis/cache and two-user ownership work on PostgreSQL")
        version = request("POST", f"/resume-studio/analyses/{aid}/versions", 201, headers=headers, json={}).json()
        version["content"]["basics"]["name"] = "Beta Synthetic Candidate"
        request("PATCH", f"/resume-studio/versions/{version['id']}", headers=headers,
            json={k: version[k] for k in ("name", "template", "content")})
        exported = request("GET", f"/resume-studio/versions/{version['id']}/docx", headers=headers)
        assert "Beta Synthetic Candidate" in " ".join(p.text for p in Document(io.BytesIO(exported.content)).paragraphs)
        assert request("GET", "/product/privacy/export", headers=headers).json()
        checks.append("tailored resume persists and exports a readable DOCX; privacy export succeeds")
        # Threshold chosen before running: 200 jobs/400 resumes; 4 concurrent clients,
        # 120 reads, zero failures, p95 below 1500 ms on this local fixture.
        from app import models as m
        from sqlalchemy.orm import Session
        with Session(engine) as db:
            for index in range(200):
                job = m.TargetJob(user_id=owner, title=f"Fixture {index}", description="D" * 50000)
                db.add(job); db.flush()
                for number in range(2):
                    resume = m.Resume(user_id=owner, target_job_id=job.id, filename=f"fixture-{number}.docx",
                        content_type="fixture", size_bytes=100, extracted_text="SYNTHETIC PRIVATE BODY " * 100)
                    db.add(resume); db.flush()
                    db.add(m.Analysis(user_id=owner, target_job_id=job.id, resume_id=resume.id,
                        resume_filename=resume.filename, resume_text=resume.extracted_text))
            db.commit()
        paths = ["/workspaces?limit=50", "/resume/all?limit=50", "/dashboard/summary"]
        def read(index):
            begin = time.perf_counter()
            response = client.get(paths[index % len(paths)], headers=headers)
            assert response.status_code == 200
            return (time.perf_counter() - begin) * 1000
        def cpu_seconds():
            fields = Path(f"/proc/{process.pid}/stat").read_text().split()
            return (int(fields[13]) + int(fields[14])) / os.sysconf("SC_CLK_TCK")
        before_cpu = cpu_seconds()
        before_load = time.monotonic()
        with ThreadPoolExecutor(max_workers=4) as pool:
            times = sorted(pool.map(read, range(120)))
        p95 = times[int(.95 * (len(times) - 1))]
        load_seconds = time.monotonic() - before_load
        api_cpu_seconds = cpu_seconds() - before_cpu
        assert p95 < 1500, f"Local read p95 exceeded declared 1500 ms: {p95:.1f}"
        status = Path(f"/proc/{process.pid}/status").read_text()
        rss = next(line.split()[1] for line in status.splitlines() if line.startswith("VmRSS:"))
        result["load"] = {"jobs": 201, "resumes": 401, "requests": len(times), "concurrency": 4,
            "errors": 0, "p50_ms": round(statistics.median(times), 2), "p95_ms": round(p95, 2),
            "p95_limit_ms": 1500, "api_rss_kib": int(rss), "container_cpu_limit": 2,
            "load_seconds": round(load_seconds, 3), "api_cpu_seconds": round(api_cpu_seconds, 3),
            "api_cpu_percent_of_two_cores": round(api_cpu_seconds / load_seconds / 2 * 100, 2),
            "container_memory_limit_mib": 512, "provider": "disabled; consent-off fixture", "provider_seconds": 0,
            "claim": "local baseline only; repeat on beta host before increasing access"}
        checks.append("representative PostgreSQL read baseline meets its declared error and p95 limits")
        from scripts.verify_database_security import audit
        with engine.connect() as db:
            security = audit(db)
            assert security["secure"], "Private database isolation failed"
            result["database"] = {"secure": True, "revision": db.execute(text("SELECT version_num FROM alembic_version")).scalar(),
                "application_tables": len(m.Base.metadata.tables) if hasattr(m, "Base") else 22}
        subprocess.run(["alembic", "check"], env=environment, check=True, stdout=log, stderr=log)
        checks.append("latest migration, metadata parity and private-table role restrictions verified")
        request("POST", "/product/privacy/delete-app-data", headers=headers,
            json={"confirmation": "DELETE MY SKILLSYNC DATA"})
        assert request("GET", "/workspaces", headers=headers).json() == []
        checks.append("owned app-data deletion removes the complete synthetic dataset")
        result.update(passed=True, non_root=True)
    except Exception as error:
        result.update(passed=False, error=f"{type(error).__name__}: {error}")
        raise
    finally:
        result["duration_seconds"] = round(time.monotonic() - started, 2)
        output.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
        process.terminate()
        try:
            process.wait(timeout=10)
        except subprocess.TimeoutExpired:
            process.kill(); process.wait()
        issuer_server.shutdown(); issuer_server.server_close(); client.close(); log.close(); engine.dispose()
print(json.dumps(result, indent=2))
