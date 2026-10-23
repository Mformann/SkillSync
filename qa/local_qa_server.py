"""Isolated real API for browser QA. Auth is stubbed; external AI is disabled."""
import io
import os
import sys
import uuid
from pathlib import Path
from datetime import date, timedelta

root = Path(__file__).resolve().parents[1]
os.environ["DATABASE_URL"] = "sqlite://"
os.environ["ENVIRONMENT"] = "development"
os.environ["GROQ_API_KEY"] = ""
os.environ["CORS_ORIGINS"] = "http://127.0.0.1:8012"
sys.path.insert(0, str(root / "backend"))

from fastapi.testclient import TestClient
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.pool import StaticPool
from sqlalchemy.orm import sessionmaker
from app.auth_utils import AuthenticatedUser, get_current_user
from app.database import Base, get_session
from app.main import create_app

engine = create_engine("sqlite:///" + (root / "qa" / ("browser-" + uuid.uuid4().hex + ".db")).as_posix(), connect_args={"check_same_thread": False})
Base.metadata.create_all(engine)
factory = sessionmaker(bind=engine)
app = create_app()
def database():
    with factory() as db:
        yield db
app.dependency_overrides[get_session] = database
app.dependency_overrides[get_current_user] = lambda: AuthenticatedUser(id="qa-browser-user", email="qa@example.com")
client = TestClient(app)
document = Document()
document.add_paragraph("Synthetic QA Candidate\nBuilt Python APIs and SQL services with automated tests.\nDelivered reliable backend applications using Git.")
content = io.BytesIO()
document.save(content)
(root / "qa" / "synthetic-resume.docx").write_bytes(content.getvalue())
table_document = Document()
table_document.add_table(rows=1, cols=1).cell(0, 0).text = document.paragraphs[0].text
table_document.save(root / "qa" / "phase04-table-resume.docx")
uploaded = client.post("/resume/upload", data={"job_role": "Backend engineer", "company": "Synthetic QA employer",
    "job_description": "We require Python, SQL, Docker and Kubernetes to build reliable production APIs and services.", "ai_consent": "false"},
    files={"file": ("synthetic-resume.docx", content.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
uploaded.raise_for_status()
aid = uploaded.json()["analysis_id"]
client.post(f"/resume/analysis/{aid}/run").raise_for_status()
client.post("/learning/plans", json={"analysis_id": aid, "hours_per_week": 5,
    "target_date": str(date.today() + timedelta(days=60)), "experience_level": "intermediate"}).raise_for_status()
client.post(f"/resume-studio/analyses/{aid}/versions", json={}).raise_for_status()

if "--phase5" in sys.argv:
    # The linked workspace deliberately falls beyond the first list page.
    for index in range(55):
        client.post("/workspaces", json={"title": f"Paging fixture {index:02d}",
            "description": "We require Python and Docker to build reliable production applications."}).raise_for_status()
    for _ in range(6):
        item = client.post("/resume/upload", data={"workspace_id": str(uploaded.json()["workspace_id"])},
            files={"file": ("history-fixture.docx", content.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        item.raise_for_status()
        client.post(f"/resume/analysis/{item.json()['analysis_id']}/run").raise_for_status()

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8011, log_level="warning")
